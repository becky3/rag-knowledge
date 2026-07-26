"""PPTX テキスト抽出・メディア削減のテスト.

仕様: docs/specs/converter.md「PPTX テキスト抽出」
関連仕様: docs/specs/infrastructure/pptx-media-reduction.md

テスト方針（aidlc-docs/plan-work/issue-799.md より転記）:
- python-pptx でフィクスチャ pptx をテスト内生成し、以下を検証する:
  - ハイブリッド形式の出力構造（`## Slide N: タイトル` / 無題スライド / `### Notes:`）
  - テーブル抽出、ノートなしスライド、空スライド
  - ppsx content-type 正規化、メディアパートを含む pptx でメディアが読まれないこと
- 削減ツール: 非破壊性（オリジナルのハッシュ不変）・削減コピー生成・削減後の抽出同一性
- Converter 経由の統合: dispatch ルーティング・壊れ pptx の errors 計上・空デッキのスキップ
"""

from __future__ import annotations

import hashlib
import io
import zipfile
from pathlib import Path

import pytest
from factories import make_converter_args
from pptx import Presentation
from pptx.util import Inches

from rag.converter.converter import (
    ConversionFailedError,
    ConversionSkippedError,
    Converter,
)
from rag.converter.pptx_extractor import (
    PptxExtractionError,
    _reduced_package_bytes,
    _title_text,
    analyze_pptx_media,
    extract_pptx,
    reduce_pptx,
)

_CONTENT_TYPES_ENTRY = "[Content_Types].xml"
_PRESENTATION_CT = b"presentationml.presentation.main+xml"
_SLIDESHOW_CT = b"presentationml.slideshow.main+xml"

# 1x1 赤ピクセルの PNG（Pillow 生成の固定バイト列を避け、テスト内で都度生成する）


def _png_bytes() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (1, 1), color=(255, 0, 0)).save(buf, format="PNG")
    return buf.getvalue()


def _new_presentation() -> Presentation:  # type: ignore[valid-type]
    return Presentation()


def _add_titled_slide(prs, title: str, body: str | None = None):  # type: ignore[no-untyped-def]
    """タイトルのみレイアウト（layouts[5]）でスライドを追加する."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    if body is not None:
        box = slide.shapes.add_textbox(
            Inches(1), Inches(2), Inches(4), Inches(1),
        )
        box.text_frame.text = body
    return slide


def _add_blank_slide(prs, body: str | None = None):  # type: ignore[no-untyped-def]
    """空白レイアウト（layouts[6]）でスライドを追加する."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if body is not None:
        box = slide.shapes.add_textbox(
            Inches(1), Inches(2), Inches(4), Inches(1),
        )
        box.text_frame.text = body
    return slide


def _save(prs, path: Path) -> Path:  # type: ignore[no-untyped-def]
    prs.save(str(path))
    return path


def _to_ppsx(src: Path, dst: Path) -> Path:
    """pptx の content-type を slideshow 形式に書き換えて ppsx を生成する."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == _CONTENT_TYPES_ENTRY:
                data = data.replace(_PRESENTATION_CT, _SLIDESHOW_CT)
            zout.writestr(info, data)
    return dst


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class TestExtractPptx:
    """extract_pptx の出力構造テスト."""

    def test_titled_slide_is_rendered_as_slide_heading_with_title(
        self, tmp_path: Path,
    ) -> None:
        prs = _new_presentation()
        _add_titled_slide(prs, "Sample Title", "Body text")
        path = _save(prs, tmp_path / "deck.pptx")

        text = extract_pptx(path)

        assert text is not None
        assert "## Slide 1: Sample Title" in text
        assert "Body text" in text

    def test_untitled_slide_heading_has_no_title(self, tmp_path: Path) -> None:
        prs = _new_presentation()
        _add_blank_slide(prs, "Only body")
        path = _save(prs, tmp_path / "deck.pptx")

        text = extract_pptx(path)

        assert text is not None
        assert "## Slide 1\n" in text
        assert "## Slide 1:" not in text

    def test_title_text_is_not_duplicated_in_body(self, tmp_path: Path) -> None:
        prs = _new_presentation()
        _add_titled_slide(prs, "Unique Heading", "Body text")
        path = _save(prs, tmp_path / "deck.pptx")

        text = extract_pptx(path)

        assert text is not None
        assert text.count("Unique Heading") == 1

    def test_slide_numbers_are_sequential(self, tmp_path: Path) -> None:
        prs = _new_presentation()
        _add_titled_slide(prs, "First")
        _add_blank_slide(prs, "second body")
        _add_titled_slide(prs, "Third")
        path = _save(prs, tmp_path / "deck.pptx")

        text = extract_pptx(path)

        assert text is not None
        assert "## Slide 1: First" in text
        assert "## Slide 2" in text
        assert "## Slide 3: Third" in text

    def test_notes_are_rendered_under_notes_heading(self, tmp_path: Path) -> None:
        prs = _new_presentation()
        slide = _add_titled_slide(prs, "With Notes", "Body")
        slide.notes_slide.notes_text_frame.text = "Speaker memo line"
        path = _save(prs, tmp_path / "deck.pptx")

        text = extract_pptx(path)

        assert text is not None
        assert "### Notes:" in text
        assert "Speaker memo line" in text
        # ノートは本文の後に配置される
        assert text.index("Body") < text.index("### Notes:")

    def test_slide_without_notes_has_no_notes_heading(self, tmp_path: Path) -> None:
        prs = _new_presentation()
        _add_titled_slide(prs, "No Notes", "Body")
        path = _save(prs, tmp_path / "deck.pptx")

        text = extract_pptx(path)

        assert text is not None
        assert "### Notes:" not in text

    def test_table_is_rendered_as_markdown_table(self, tmp_path: Path) -> None:
        prs = _new_presentation()
        slide = _add_blank_slide(prs)
        table = slide.shapes.add_table(
            2, 2, Inches(1), Inches(1), Inches(4), Inches(1),
        ).table
        table.cell(0, 0).text = "Header A"
        table.cell(0, 1).text = "Header B"
        table.cell(1, 0).text = "Cell 1"
        table.cell(1, 1).text = "Cell 2"
        path = _save(prs, tmp_path / "deck.pptx")

        text = extract_pptx(path)

        assert text is not None
        assert "| Header A | Header B |" in text
        assert "| --- | --- |" in text
        assert "| Cell 1 | Cell 2 |" in text

    def test_group_shape_texts_are_extracted_recursively(
        self, tmp_path: Path,
    ) -> None:
        prs = _new_presentation()
        slide = _add_blank_slide(prs)
        group = slide.shapes.add_group_shape()
        box = group.shapes.add_textbox(
            Inches(1), Inches(1), Inches(3), Inches(1),
        )
        box.text_frame.text = "Grouped text"
        path = _save(prs, tmp_path / "deck.pptx")

        text = extract_pptx(path)

        assert text is not None
        assert "Grouped text" in text

    def test_empty_deck_returns_none(self, tmp_path: Path) -> None:
        prs = _new_presentation()
        _add_blank_slide(prs)
        path = _save(prs, tmp_path / "deck.pptx")

        assert extract_pptx(path) is None

    def test_deck_without_slides_returns_none(self, tmp_path: Path) -> None:
        prs = _new_presentation()
        path = _save(prs, tmp_path / "deck.pptx")

        assert extract_pptx(path) is None

    def test_broken_zip_raises_extraction_error(self, tmp_path: Path) -> None:
        path = tmp_path / "broken.pptx"
        path.write_bytes(b"this is not a zip file")

        with pytest.raises(PptxExtractionError):
            extract_pptx(path)

    def test_zero_byte_file_raises_extraction_error(self, tmp_path: Path) -> None:
        path = tmp_path / "empty.pptx"
        path.write_bytes(b"")

        with pytest.raises(PptxExtractionError):
            extract_pptx(path)

    def test_ppsx_with_slideshow_content_type_is_extracted(
        self, tmp_path: Path,
    ) -> None:
        prs = _new_presentation()
        _add_titled_slide(prs, "Slideshow Deck", "Body")
        pptx_path = _save(prs, tmp_path / "deck.pptx")
        ppsx_path = _to_ppsx(pptx_path, tmp_path / "deck.ppsx")

        text = extract_pptx(ppsx_path)

        assert text is not None
        assert "## Slide 1: Slideshow Deck" in text

    def test_media_parts_are_emptied_in_reduced_package(
        self, tmp_path: Path,
    ) -> None:
        prs = _new_presentation()
        slide = _add_titled_slide(prs, "With Image", "Body")
        slide.shapes.add_picture(
            io.BytesIO(_png_bytes()), Inches(1), Inches(3),
        )
        path = _save(prs, tmp_path / "deck.pptx")

        package = _reduced_package_bytes(path)
        with zipfile.ZipFile(io.BytesIO(package)) as zf:
            media_entries = [
                info for info in zf.infolist()
                if info.filename.startswith("ppt/media/")
            ]
        assert media_entries
        assert all(info.file_size == 0 for info in media_entries)

        # メディアを空置換してもテキスト抽出は成立する
        text = extract_pptx(path)
        assert text is not None
        assert "## Slide 1: With Image" in text


class TestTitleText:
    """_title_text のガードテスト.

    タイトル位置のプレースホルダーが画像に置換された実データ
    （PlaceholderPicture が返り `.text` を持たない）で AttributeError に
    ならないことを検証する（L3 QA で検出された実ケース）。
    """

    def test_none_title_returns_empty(self) -> None:
        assert _title_text(None) == ""

    def test_non_text_placeholder_returns_empty(self) -> None:
        class _FakePlaceholderPicture:
            """text / text_frame を持たない図形のスタブ."""

            has_text_frame = False

        assert _title_text(_FakePlaceholderPicture()) == ""  # type: ignore[arg-type]

    def test_text_title_shape_returns_text(self) -> None:
        prs = _new_presentation()
        slide = _add_titled_slide(prs, "Real Title")
        assert _title_text(slide.shapes.title) == "Real Title"


class TestAnalyzePptxMedia:
    """analyze_pptx_media のレポートテスト."""

    def test_media_report_counts_media_parts(self, tmp_path: Path) -> None:
        prs = _new_presentation()
        slide = _add_blank_slide(prs, "body")
        slide.shapes.add_picture(
            io.BytesIO(_png_bytes()), Inches(1), Inches(3),
        )
        path = _save(prs, tmp_path / "deck.pptx")

        report = analyze_pptx_media(path)

        assert report.total_bytes == path.stat().st_size
        assert report.media_count == 1
        assert report.media_ext_counts == {"png": 1}
        assert 0 <= report.reduced_estimate_bytes <= report.total_bytes

    def test_deck_without_media_reports_zero(self, tmp_path: Path) -> None:
        prs = _new_presentation()
        _add_blank_slide(prs, "body")
        path = _save(prs, tmp_path / "deck.pptx")

        report = analyze_pptx_media(path)

        assert report.media_count == 0
        assert report.media_bytes == 0

    def test_broken_zip_raises_extraction_error(self, tmp_path: Path) -> None:
        path = tmp_path / "broken.pptx"
        path.write_bytes(b"garbage")

        with pytest.raises(PptxExtractionError):
            analyze_pptx_media(path)


class TestReducePptx:
    """reduce_pptx の非破壊性・削減効果テスト."""

    def _deck_with_picture(self, tmp_path: Path) -> Path:
        prs = _new_presentation()
        slide = _add_titled_slide(prs, "Reduce Target", "Body text")
        slide.shapes.add_picture(
            io.BytesIO(_png_bytes()), Inches(1), Inches(3),
        )
        return _save(prs, tmp_path / "deck.pptx")

    def test_reduce_is_non_destructive(self, tmp_path: Path) -> None:
        src = self._deck_with_picture(tmp_path)
        hash_before = _sha256(src)

        reduce_pptx(src, tmp_path / "reduced.pptx")

        assert _sha256(src) == hash_before

    def test_reduced_copy_has_empty_media_entries(self, tmp_path: Path) -> None:
        src = self._deck_with_picture(tmp_path)
        dst = tmp_path / "reduced.pptx"

        reduce_pptx(src, dst)

        with zipfile.ZipFile(dst) as zf:
            media_entries = [
                info for info in zf.infolist()
                if info.filename.startswith("ppt/media/")
            ]
        assert media_entries
        assert all(info.file_size == 0 for info in media_entries)

    def test_reduced_copy_extraction_is_identical(self, tmp_path: Path) -> None:
        src = self._deck_with_picture(tmp_path)
        dst = tmp_path / "reduced.pptx"

        reduce_pptx(src, dst)

        assert extract_pptx(dst) == extract_pptx(src)

    def test_reduce_empties_embeddings_and_fonts_entries(
        self, tmp_path: Path,
    ) -> None:
        """埋め込みオブジェクト・フォントのパートも空置換される（削減対象パート）."""
        src = self._deck_with_picture(tmp_path)
        # 削減対象は名前ベースのため、embeddings / fonts エントリを直接注入して検証する
        with zipfile.ZipFile(src, "a") as zf:
            zf.writestr("ppt/embeddings/oleObject1.bin", b"x" * 1024)
            zf.writestr("ppt/fonts/font1.fntdata", b"y" * 2048)
        dst = tmp_path / "reduced.pptx"

        reduce_pptx(src, dst)

        with zipfile.ZipFile(dst) as zf:
            emptied = {
                info.filename: info.file_size
                for info in zf.infolist()
                if info.filename.startswith(("ppt/embeddings/", "ppt/fonts/"))
            }
        assert emptied == {
            "ppt/embeddings/oleObject1.bin": 0,
            "ppt/fonts/font1.fntdata": 0,
        }

    def test_reduce_keeps_ppsx_content_type_unchanged(self, tmp_path: Path) -> None:
        prs = _new_presentation()
        _add_titled_slide(prs, "Slideshow", "Body")
        pptx_path = _save(prs, tmp_path / "deck.pptx")
        ppsx_path = _to_ppsx(pptx_path, tmp_path / "deck.ppsx")
        dst = tmp_path / "reduced.ppsx"

        reduce_pptx(ppsx_path, dst)

        with zipfile.ZipFile(dst) as zf:
            content_types = zf.read(_CONTENT_TYPES_ENTRY)
        assert _SLIDESHOW_CT in content_types

    def test_reduce_broken_zip_raises_and_leaves_no_output(
        self, tmp_path: Path,
    ) -> None:
        src = tmp_path / "broken.pptx"
        src.write_bytes(b"garbage")
        dst = tmp_path / "reduced.pptx"

        with pytest.raises(PptxExtractionError):
            reduce_pptx(src, dst)

        assert not dst.exists()


def _setup_stores(tmp_path: Path) -> tuple[Path, Path]:
    """テスト用の source_store / converted_store ディレクトリを返す."""
    source_dir = tmp_path / "source_store"
    converted_dir = tmp_path / "converted_store"
    source_dir.mkdir()
    return source_dir, converted_dir


class TestConverterPptxIntegration:
    """Converter.convert / convert_batch 経由の pptx 変換テスト.

    仕様: docs/specs/converter.md「PPTX テキスト抽出」「スキップと失敗の区別」
    """

    def _place_deck(
        self, source_dir: Path, rel_path: str, *, title: str, body: str,
    ) -> None:
        prs = _new_presentation()
        _add_titled_slide(prs, title, body)
        dest = source_dir / rel_path
        dest.parent.mkdir(parents=True, exist_ok=True)
        _save(prs, dest)

    def test_pptx_is_converted_to_md_via_dispatch(self, tmp_path: Path) -> None:
        source_dir, converted_dir = _setup_stores(tmp_path)
        self._place_deck(
            source_dir, "local/docs/talk.pptx",
            title="Talk Title", body="Talk body",
        )
        converter = Converter(**make_converter_args())

        result_path = converter.convert(
            "local/docs/talk.pptx", source_dir, converted_dir,
        )

        assert result_path == converted_dir / "local/docs/talk.md"
        content = result_path.read_text(encoding="utf-8")
        assert "## Slide 1: Talk Title" in content
        assert "Talk body" in content

    def test_ppsx_is_converted_to_md_via_dispatch(self, tmp_path: Path) -> None:
        source_dir, converted_dir = _setup_stores(tmp_path)
        prs = _new_presentation()
        _add_titled_slide(prs, "Slideshow Title", "Body")
        pptx_path = _save(prs, tmp_path / "work.pptx")
        dest = source_dir / "local/docs/show.ppsx"
        dest.parent.mkdir(parents=True, exist_ok=True)
        _to_ppsx(pptx_path, dest)
        converter = Converter(**make_converter_args())

        result_path = converter.convert(
            "local/docs/show.ppsx", source_dir, converted_dir,
        )

        assert result_path == converted_dir / "local/docs/show.md"
        assert "## Slide 1: Slideshow Title" in result_path.read_text(
            encoding="utf-8",
        )

    def test_broken_pptx_raises_conversion_failed_error(
        self, tmp_path: Path,
    ) -> None:
        source_dir, converted_dir = _setup_stores(tmp_path)
        dest = source_dir / "local/docs/broken.pptx"
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(b"this is not a zip file")
        converter = Converter(**make_converter_args())

        with pytest.raises(ConversionFailedError):
            converter.convert("local/docs/broken.pptx", source_dir, converted_dir)

    def test_broken_pptx_is_counted_as_error_in_convert_batch(
        self, tmp_path: Path,
    ) -> None:
        source_dir, converted_dir = _setup_stores(tmp_path)
        broken = source_dir / "local/docs/broken.pptx"
        broken.parent.mkdir(parents=True, exist_ok=True)
        broken_bytes = b"this is not a zip file"
        broken.write_bytes(broken_bytes)
        converter = Converter(**make_converter_args())

        result = converter.convert_batch(
            ["local/docs/broken.pptx"], source_dir, converted_dir,
        )

        assert result.errors == 1
        entry = result.error_files[0]
        assert entry["path"] == "local/docs/broken.pptx"
        assert entry["size_bytes"] == len(broken_bytes)

    def test_empty_deck_raises_conversion_skipped_error(
        self, tmp_path: Path,
    ) -> None:
        source_dir, converted_dir = _setup_stores(tmp_path)
        prs = _new_presentation()
        _add_blank_slide(prs)
        dest = source_dir / "local/docs/empty-deck.pptx"
        dest.parent.mkdir(parents=True, exist_ok=True)
        _save(prs, dest)
        converter = Converter(**make_converter_args())

        with pytest.raises(ConversionSkippedError):
            converter.convert(
                "local/docs/empty-deck.pptx", source_dir, converted_dir,
            )
        assert not (converted_dir / "local/docs/empty-deck.md").exists()
