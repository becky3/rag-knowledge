"""reduce-pptx CLI コマンドのテスト.

仕様: docs/specs/infrastructure/pptx-media-reduction.md

テスト方針（aidlc-docs/plan-work/issue-799.md より転記）:
- レポートのみモード・削減実行・既存出力のスキップ・出力先未指定エラーを検証する
- --alongside 別名出力（<元名>.reduced.<拡張子>）・--output-dir との排他・
  .reduced サフィックス付きファイルの対象除外を検証する
- フィクスチャ pptx は python-pptx でテスト内生成する
"""

from __future__ import annotations

import argparse
import io
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from rag.cli import run_reduce_pptx


def _make_deck(path: Path, *, with_picture: bool = True) -> Path:
    from PIL import Image

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "CLI Test Deck"
    if with_picture:
        buf = io.BytesIO()
        Image.new("RGB", (1, 1), color=(0, 0, 255)).save(buf, format="PNG")
        buf.seek(0)
        slide.shapes.add_picture(buf, Inches(1), Inches(3))
    prs.save(str(path))
    return path


def _args(
    paths: list[str],
    *,
    output_dir: str | None = None,
    alongside: bool = False,
    report_only: bool = False,
) -> argparse.Namespace:
    return argparse.Namespace(
        paths=paths,
        output_dir=output_dir,
        alongside=alongside,
        report_only=report_only,
    )


class TestRunReducePptx:
    """run_reduce_pptx の振る舞いテスト."""

    def test_report_only_shows_media_summary_without_output(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str],
    ) -> None:
        deck = _make_deck(tmp_path / "deck.pptx")

        run_reduce_pptx(_args([str(deck)], report_only=True))

        out = capsys.readouterr().out
        assert "削減対象: 1 件" in out
        assert "レポート完了: 対象 1 件 / エラー 0 件" in out
        assert not list(tmp_path.glob("*/deck.pptx"))

    def test_reduction_creates_copy_in_output_dir(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str],
    ) -> None:
        deck = _make_deck(tmp_path / "deck.pptx")
        out_dir = tmp_path / "reduced"

        run_reduce_pptx(_args([str(deck)], output_dir=str(out_dir)))

        out = capsys.readouterr().out
        assert (out_dir / "deck.pptx").exists()
        assert "削減完了: 生成 1 件 / スキップ 0 件 / エラー 0 件" in out

    def test_directory_input_collects_pptx_recursively(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str],
    ) -> None:
        src_dir = tmp_path / "corpus"
        (src_dir / "nested").mkdir(parents=True)
        _make_deck(src_dir / "a.pptx")
        _make_deck(src_dir / "nested" / "b.pptx")
        (src_dir / "note.txt").write_text("not pptx", encoding="utf-8")
        out_dir = tmp_path / "reduced"

        run_reduce_pptx(_args([str(src_dir)], output_dir=str(out_dir)))

        out = capsys.readouterr().out
        assert (out_dir / "a.pptx").exists()
        assert (out_dir / "b.pptx").exists()
        assert "生成 2 件" in out

    def test_existing_output_is_skipped_not_overwritten(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str],
    ) -> None:
        deck = _make_deck(tmp_path / "deck.pptx")
        out_dir = tmp_path / "reduced"
        out_dir.mkdir()
        existing = out_dir / "deck.pptx"
        existing.write_bytes(b"existing content")

        run_reduce_pptx(_args([str(deck)], output_dir=str(out_dir)))

        out = capsys.readouterr().out
        assert existing.read_bytes() == b"existing content"
        assert "スキップ 1 件" in out

    def test_missing_output_dir_without_report_only_exits(
        self, tmp_path: Path,
    ) -> None:
        deck = _make_deck(tmp_path / "deck.pptx")

        with pytest.raises(SystemExit) as exc_info:
            run_reduce_pptx(_args([str(deck)]))

        assert exc_info.value.code == 1

    def test_output_dir_and_alongside_are_mutually_exclusive(
        self, tmp_path: Path,
    ) -> None:
        deck = _make_deck(tmp_path / "deck.pptx")

        with pytest.raises(SystemExit) as exc_info:
            run_reduce_pptx(_args(
                [str(deck)], output_dir=str(tmp_path / "out"), alongside=True,
            ))

        assert exc_info.value.code == 1

    def test_alongside_creates_sibling_reduced_copy(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str],
    ) -> None:
        deck = _make_deck(tmp_path / "deck.pptx")

        run_reduce_pptx(_args([str(deck)], alongside=True))

        out = capsys.readouterr().out
        reduced = tmp_path / "deck.reduced.pptx"
        assert reduced.exists()
        assert reduced.stat().st_size < deck.stat().st_size
        assert "生成 1 件" in out

    def test_alongside_skips_existing_reduced_copy(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str],
    ) -> None:
        deck = _make_deck(tmp_path / "deck.pptx")
        existing = tmp_path / "deck.reduced.pptx"
        existing.write_bytes(b"existing content")

        run_reduce_pptx(_args([str(deck)], alongside=True))

        out = capsys.readouterr().out
        assert existing.read_bytes() == b"existing content"
        assert "スキップ 1 件" in out

    def test_already_reduced_files_are_excluded_from_targets(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str],
    ) -> None:
        src_dir = tmp_path / "corpus"
        src_dir.mkdir()
        _make_deck(src_dir / "deck.pptx")
        _make_deck(src_dir / "deck.reduced.pptx")

        run_reduce_pptx(_args([str(src_dir)], report_only=True))

        out = capsys.readouterr().out
        assert "レポート完了: 対象 1 件" in out

    def test_broken_file_is_counted_as_error_and_processing_continues(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str],
    ) -> None:
        broken = tmp_path / "broken.pptx"
        broken.write_bytes(b"garbage")
        deck = _make_deck(tmp_path / "deck.pptx")
        out_dir = tmp_path / "reduced"

        run_reduce_pptx(_args([str(broken), str(deck)], output_dir=str(out_dir)))

        out = capsys.readouterr().out
        assert (out_dir / "deck.pptx").exists()
        assert "生成 1 件" in out
        assert "エラー 1 件" in out

    def test_nonexistent_path_is_reported_as_error(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str],
    ) -> None:
        deck = _make_deck(tmp_path / "deck.pptx")
        out_dir = tmp_path / "reduced"

        run_reduce_pptx(_args(
            [str(tmp_path / "missing.pptx"), str(deck)],
            output_dir=str(out_dir),
        ))

        out = capsys.readouterr().out
        assert "エラー 1 件" in out
        assert (out_dir / "deck.pptx").exists()

    def test_all_files_error_exits_nonzero(self, tmp_path: Path) -> None:
        broken = tmp_path / "broken.pptx"
        broken.write_bytes(b"garbage")
        out_dir = tmp_path / "reduced"

        with pytest.raises(SystemExit) as exc_info:
            run_reduce_pptx(_args([str(broken)], output_dir=str(out_dir)))

        assert exc_info.value.code == 1

    def test_report_only_all_files_error_exits_nonzero(
        self, tmp_path: Path,
    ) -> None:
        broken = tmp_path / "broken.pptx"
        broken.write_bytes(b"garbage")

        with pytest.raises(SystemExit) as exc_info:
            run_reduce_pptx(_args([str(broken)], report_only=True))

        assert exc_info.value.code == 1

    def test_same_name_collision_within_run_skips_second_file(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str],
    ) -> None:
        dir_a = tmp_path / "a"
        dir_b = tmp_path / "b"
        dir_a.mkdir()
        dir_b.mkdir()
        deck_a = _make_deck(dir_a / "deck.pptx")
        deck_b = _make_deck(dir_b / "deck.pptx")
        out_dir = tmp_path / "reduced"

        run_reduce_pptx(_args(
            [str(deck_a), str(deck_b)], output_dir=str(out_dir),
        ))

        out = capsys.readouterr().out
        assert "同一実行内で出力名が衝突するためスキップ" in out
        assert "生成 1 件" in out
        assert "スキップ 1 件" in out
