"""PPTX テキスト抽出モジュール.

仕様: docs/specs/converter.md
関連仕様: docs/specs/infrastructure/pptx-media-reduction.md

pptx/ppsx（zip コンテナ）からスライド本文・テーブル・スピーカーノートを抽出し、
Markdown 形式で出力する。削減対象パート（ppt/media/* / ppt/embeddings/* /
ppt/fonts/*）の実体は読み込まず、空データに置換した in-memory パッケージを
python-pptx に渡すことで、動画埋め込みで GB 級のファイルでも
テキスト XML 分の処理量に抑える。

メディア実体の空置換ロジックは pptx メディア削減ツール（CLI reduce-pptx）と共有する。
"""

from __future__ import annotations

import io
import logging
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pptx.shapes.autoshape import Shape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape

if TYPE_CHECKING:
    from collections.abc import Iterable

    from pptx.shapes.base import BaseShape
    from pptx.table import Table

logger = logging.getLogger(__name__)

# pptx/ppsx の拡張子（converter のディスパッチからも参照される public 定数）
PPTX_EXTENSIONS: frozenset[str] = frozenset({".pptx", ".ppsx"})

# 削減コピーの別名出力サフィックス（reduce-pptx --alongside が使用。
# このサフィックスを持つファイルは削減対象の収集から除外される）
REDUCED_STEM_SUFFIX = ".reduced"

# 削減対象パートの zip エントリプレフィックス
# （メディア実体・埋め込みオブジェクト・埋め込みフォント。テキスト抽出には使われない）
_STRIP_PREFIXES: tuple[str, str, str] = (
    "ppt/media/",
    "ppt/embeddings/",
    "ppt/fonts/",
)

_CONTENT_TYPES_ENTRY = "[Content_Types].xml"

# ppsx（スライドショー形式）の main part content-type。
# python-pptx はこの content-type を presentation として認識しないため、
# 変換時に presentation 形式へ正規化する
_SLIDESHOW_MAIN_CT = (
    "application/vnd.openxmlformats-officedocument"
    ".presentationml.slideshow.main+xml"
)
_PRESENTATION_MAIN_CT = (
    "application/vnd.openxmlformats-officedocument"
    ".presentationml.presentation.main+xml"
)


class PptxExtractionError(Exception):
    """pptx/ppsx がパッケージとして開けない場合の例外（壊れ zip 等）."""


@dataclass(frozen=True)
class PptxMediaReport:
    """pptx の削減対象パート（メディア・埋め込みオブジェクト・フォント）の占有量レポート.

    media_bytes は zip 内の圧縮後占有量（ディスク上のファイルサイズへの寄与分）。
    非圧縮サイズを使うと、圧縮の効くメディア（EMF 等）で削減後推定が
    ファイルサイズを超えて負値になるため採らない。
    """

    total_bytes: int
    media_count: int
    media_bytes: int
    media_ext_counts: dict[str, int]

    @property
    def reduced_estimate_bytes(self) -> int:
        """削減後の推定サイズ（メディアの zip 内占有量を除いた概算）."""
        return max(self.total_bytes - self.media_bytes, 0)


def analyze_pptx_media(path: Path) -> PptxMediaReport:
    """pptx/ppsx の削減対象パートの占有量を解析する.

    zip セントラルディレクトリのみを読むため、GB 級ファイルでも軽量。

    Args:
        path: 対象ファイルパス

    Returns:
        メディア占有量レポート

    Raises:
        PptxExtractionError: zip として開けない場合
    """
    try:
        total_bytes = path.stat().st_size
        with zipfile.ZipFile(path) as zf:
            media_infos = [
                info for info in zf.infolist()
                if info.filename.startswith(_STRIP_PREFIXES)
            ]
    except (zipfile.BadZipFile, OSError) as e:
        msg = f"pptx を zip として開けません: {path.name} ({e})"
        raise PptxExtractionError(msg) from e

    ext_counts: dict[str, int] = {}
    for info in media_infos:
        ext = Path(info.filename).suffix.lower().lstrip(".") or "(none)"
        ext_counts[ext] = ext_counts.get(ext, 0) + 1

    return PptxMediaReport(
        total_bytes=total_bytes,
        media_count=len(media_infos),
        media_bytes=sum(info.compress_size for info in media_infos),
        media_ext_counts=ext_counts,
    )


def _write_reduced_entries(
    zin: zipfile.ZipFile,
    zout: zipfile.ZipFile,
    *,
    normalize_content_type: bool,
) -> None:
    """zip エントリをコピーし、削減対象パートの実体のみ空データに置換する.

    エントリ自体・リレーションは維持するため、削減後も正当な pptx として開ける。

    Args:
        zin: 入力 zip（読み取り専用）
        zout: 出力 zip
        normalize_content_type: ppsx の main part content-type を
            presentation 形式へ正規化するか（変換時のみ True）
    """
    for info in zin.infolist():
        name = info.filename
        if name.startswith(_STRIP_PREFIXES):
            zout.writestr(name, b"")
        elif normalize_content_type and name == _CONTENT_TYPES_ENTRY:
            xml = zin.read(name).decode("utf-8")
            zout.writestr(
                name,
                xml.replace(_SLIDESHOW_MAIN_CT, _PRESENTATION_MAIN_CT),
            )
        else:
            zout.writestr(info, zin.read(name))


def reduce_pptx(src: Path, dst: Path) -> None:
    """削減対象パート（メディア・埋め込みオブジェクト・フォント）の実体を空置換した削減コピーを生成する（非破壊）.

    入力は読み取り専用で開き、一切変更しない。content-type は正規化しない
    （ppsx は ppsx のまま出力する。サイズ削減のみが本関数の責務）。

    Args:
        src: 入力 pptx/ppsx のパス
        dst: 削減コピーの出力先パス

    Raises:
        PptxExtractionError: 入力が zip として開けない場合
    """
    try:
        with (
            zipfile.ZipFile(src) as zin,
            zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout,
        ):
            _write_reduced_entries(zin, zout, normalize_content_type=False)
    except zipfile.BadZipFile as e:
        dst.unlink(missing_ok=True)
        msg = f"pptx を zip として開けません: {src.name} ({e})"
        raise PptxExtractionError(msg) from e
    except OSError as e:
        # 入力 open / 出力 write のどちらの失敗か特定できないため、入出力中立の文言にする
        dst.unlink(missing_ok=True)
        msg = f"pptx の削減コピー生成に失敗しました: {src.name} -> {dst} ({e})"
        raise PptxExtractionError(msg) from e


def _reduced_package_bytes(path: Path) -> bytes:
    """削減対象パート空置換 + content-type 正規化済みの in-memory パッケージを構築する.

    Args:
        path: 対象 pptx/ppsx のパス

    Returns:
        正規化済みパッケージのバイト列（メディア実体を含まないため小さい）

    Raises:
        PptxExtractionError: zip として開けない場合
    """
    try:
        with zipfile.ZipFile(path) as zin:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
                _write_reduced_entries(zin, zout, normalize_content_type=True)
            return buf.getvalue()
    except (zipfile.BadZipFile, OSError) as e:
        msg = f"pptx を zip として開けません: {path.name} ({e})"
        raise PptxExtractionError(msg) from e


def _clean_text(text: str) -> str:
    """テキストフレームのテキストを整形する（垂直タブ→改行、前後空白除去）."""
    return text.replace("\x0b", "\n").strip()


def _clean_cell(text: str) -> str:
    """テーブルセルのテキストを Markdown セル向けに整形する."""
    return text.replace("\n", " ").replace("\x0b", " ").replace("|", "\\|").strip()


def _table_to_markdown(table: Table) -> str:
    """テーブルを Markdown テーブルに変換する（先頭行をヘッダーとして扱う）."""
    rows = [[_clean_cell(cell.text) for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header, *body = rows
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _title_text(title_shape: BaseShape | None) -> str:
    """タイトルプレースホルダーのテキストを取得する.

    タイトル位置のプレースホルダーが画像に置換されたデッキでは、
    テキストを持たない図形（PlaceholderPicture 等）が返るため、
    テキストフレームを持つ図形のみからテキストを取得する。

    Args:
        title_shape: タイトルプレースホルダー（存在しない場合は None）

    Returns:
        タイトルテキスト（取得できない場合は空文字列）
    """
    if isinstance(title_shape, Shape) and title_shape.has_text_frame:
        return _clean_text(title_shape.text_frame.text)
    return ""


def _collect_shape_texts(
    shapes: Iterable[BaseShape],
    exclude_shape_id: int | None,
) -> list[str]:
    """図形からテキストブロックを抽出する（スライド XML の定義順、グループは再帰展開）.

    Args:
        shapes: 対象の図形群
        exclude_shape_id: 除外する図形 ID（タイトルプレースホルダー。見出しに出力済みのため）

    Returns:
        テキストブロックのリスト（空ブロックは含まない）
    """
    texts: list[str] = []
    for shape in shapes:
        if exclude_shape_id is not None and shape.shape_id == exclude_shape_id:
            continue
        if isinstance(shape, GroupShape):
            texts.extend(_collect_shape_texts(shape.shapes, None))
        elif isinstance(shape, GraphicFrame):
            if shape.has_table:
                table_md = _table_to_markdown(shape.table)
                if table_md:
                    texts.append(table_md)
        elif isinstance(shape, Shape) and shape.has_text_frame:
            text = _clean_text(shape.text_frame.text)
            if text:
                texts.append(text)
    return texts


def extract_pptx(path: Path) -> str | None:
    """pptx/ppsx からテキストを抽出し、Markdown 形式で返す.

    出力構造（ハイブリッド形式）:
    - スライド区切りは ``## Slide N: <タイトル>`` 見出し（無題スライドは ``## Slide N``）
    - タイトルは見出しに含め、本文には重複出力しない
    - ノートが存在するスライドのみ、本文直後に ``### Notes:`` 見出しで付加する

    Args:
        path: 対象 pptx/ppsx のパス

    Returns:
        抽出テキスト（Markdown）。デッキ全体でテキストが空の場合は None

    Raises:
        PptxExtractionError: パッケージとして開けない場合（壊れ zip 等）
    """
    package_bytes = _reduced_package_bytes(path)
    try:
        presentation = Presentation(io.BytesIO(package_bytes))
    except (PackageNotFoundError, KeyError) as e:
        msg = f"pptx をパッケージとして開けません: {path.name} ({e})"
        raise PptxExtractionError(msg) from e

    sections: list[str] = []
    has_content = False
    for idx, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title_text = _title_text(title_shape)
        heading = f"## Slide {idx}: {title_text}" if title_text else f"## Slide {idx}"

        # 除外の目的は「見出しに出力済みテキストの重複防止」のため、
        # タイトルテキストを取得できなかった場合は除外しない
        # （タイトル位置がテーブル等に置換されたデッキでの silent なテキスト損失を防ぐ）
        body_parts = _collect_shape_texts(
            slide.shapes,
            title_shape.shape_id if title_shape is not None and title_text else None,
        )

        notes_text = ""
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                notes_text = _clean_text(notes_frame.text)

        if title_text or body_parts or notes_text:
            has_content = True

        section_parts = [heading, *body_parts]
        if notes_text:
            section_parts.append("### Notes:")
            section_parts.append(notes_text)
        sections.append("\n\n".join(section_parts))

    if not has_content:
        logger.warning("No text found in pptx: %s", path.name)
        return None

    return "\n\n".join(sections) + "\n"
